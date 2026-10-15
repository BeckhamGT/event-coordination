"""Read only disclosed sources; preserve observations before interpreting them."""
import csv
import hashlib
import io
import json
import re
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pypdf import PdfReader


class InputError(ValueError):
    pass


def now():
    return datetime.now(timezone.utc).isoformat()


def digest(data):
    return "sha256:" + hashlib.sha256(data).hexdigest()


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def integer(value, field="number"):
    try:
        number = Decimal(str(value))
        if not number.is_finite() or number < 0 or number != number.to_integral_value():
            raise ValueError()
        return int(number)
    except (InvalidOperation, ValueError, TypeError):
        raise InputError("Invalid non-negative integer for " + field)


def parse_table(path, required_fields, key_field):
    path = Path(path)
    if path.suffix == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=True)
        if len(book.sheetnames) != 1:
            raise InputError("Ambiguous workbook tabs")
        rows = list(book.active.values)
        book.close()
    else:
        rows = list(csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig"))))
    if not rows:
        raise InputError("Empty table")
    headers = [str(x).strip() if x is not None else "" for x in rows[0]]
    named = [x for x in headers if x]
    if len(named) != len(set(named)):
        raise InputError("Duplicate header")
    missing = set(required_fields) - set(headers)
    if missing:
        raise InputError("Missing required fields: " + ",".join(sorted(missing)))
    output, identities = [], set()
    for row in rows[1:]:
        if not any(x is not None and str(x).strip() for x in row):
            continue
        record = {}
        for index, name in enumerate(headers):
            if not name:
                continue
            value = row[index] if index < len(row) else None
            if isinstance(value, datetime):
                value = value.isoformat()
            record[name] = str(value).strip() if value is not None else ""
        identity = record.get(key_field, "")
        if not identity or identity in identities:
            raise InputError("Missing or duplicate identity: " + identity)
        identities.add(identity)
        for key in required_fields:
            # Empty notes are permissible, missing business quantities are not.
            if not record.get(key) and key not in ("notes",):
                raise InputError("Missing value for " + identity + ":" + key)
            if key in ("value", "capacity", "quote_amount_twd", "planned_amount_twd", "approval_limit_twd", "committed_amount_twd"):
                integer(record[key], identity + ":" + key)
        output.append(record)
    return output


def _request(url, data=None):
    request = urllib.request.Request(url, data=data, headers={"User-Agent": "ProjectC-EvidenceReader/1.0", "Content-Type": "application/json"})
    with urllib.request.urlopen(request, timeout=25) as response:
        body = response.read(10 * 1024 * 1024 + 1)
        if len(body) > 10 * 1024 * 1024:
            raise InputError("Source exceeds10MB capture limit")
        return body, {key: response.headers.get(key) for key in ("Content-Type", "Last-Modified", "ETag", "Date")}, response.status


def _notion_text(body, requested_id=None):
    data = json.loads(body)
    blocks = data.get("recordMap", {}).get("block", {})
    if not blocks:
        raise InputError("Notion response has no accessible block records")
    def unwrap(value):
        while isinstance(value, dict) and "type" not in value and isinstance(value.get("value"), dict):
            value = value["value"]
        return value
    values = {key: unwrap(value) for key, value in blocks.items()}
    seen, text = set(), []

    def visit(key):
        if key in seen:
            return
        if key not in values:
            raise InputError("Incomplete Notion child block: " + key)
        seen.add(key)
        value = values[key]
        title = value.get("properties", {}).get("title", [])
        if title:
            text.append("".join(str(part[0]) for part in title if part))
        for child in value.get("content", []):
            visit(child)
    page = requested_id if requested_id in values else next((key for key, value in values.items() if value.get("type") == "page"), None)
    if page:
        visit(page)
    else:
        for key in values:
            visit(key)
    result = "\n".join(text)
    if "EVENT-BRIEF" not in result:
        raise InputError("Notion page identity/content could not be verified")
    return result


def capture(spec, evidence, visual_review, bundle=None):
    """One source record, with each actual attempted transport recorded separately."""
    evidence.mkdir(parents=True, exist_ok=True)
    name, sid = spec["name"], spec["id"]
    destination = evidence / spec["filename"]
    source = {"id": sid, "summary": name, "evidence_ids": [sid], "source_role": spec["role"], "locator": spec["url"], "retrieved_at": now(), "retrieval_status": "unavailable", "content_type": "application/octet-stream", "version_metadata": {}, "content_hash": None, "local_reference": None, "observations": [], "attempts": []}
    body = None
    if bundle is not None:
        original_registry = Path(bundle) / "source-registry.json"
        prior = json.loads(original_registry.read_text()) if original_registry.exists() else {"sources": []}
        old = next((x for x in prior["sources"] if x["id"] == sid), None)
        try:
            body = (Path(bundle) / spec["filename"]).read_bytes()
            if old is None or old.get("content_hash") != digest(body):
                raise InputError("Observed-source bundle lacks matching source identity/hash")
            source.update({"retrieved_at": old["retrieved_at"], "content_type": old["content_type"], "version_metadata": old.get("version_metadata", {})})
            source["reuse_checked_at"] = now()
            source["mode"] = "explicit-observed-source-audit;not a new remote retrieval"
            source["attempts"].append({"method": "reopen-hash-verify-observed-source", "observed_at": now(), "status": "retrieved"})
        except Exception as error:
            body = None
            source["attempts"].append({"method": "observed-source-audit", "observed_at": now(), "status": "unavailable", "error": str(error)})
    else:
        urls = [(spec["url"], None)]
        if spec["kind"] == "xlsx":
            urls = [(spec["url"] + "/export?format=xlsx", None)]
        elif spec["kind"] == "notion":
            payload = json.dumps({"pageId": spec["page_id"], "limit": 100, "cursor": {"stack": []}, "chunkNumber": 0, "verticalColumns": False}).encode()
            urls = [(spec["public_url"], None), ("https://www.notion.so/api/v3/loadPageChunk", payload), ("https://private-pecorino-70e.notion.site/api/v3/loadPageChunk", payload)]
        for attempt, (url, payload) in enumerate(urls, 1):
            observed = now()
            try:
                raw, headers, status = _request(url, payload)
                # Keep only the permitted business-text extract for Notion, not
                # account/user maps or login-page tokens. Successful binary/HTML
                # source bytes are preserved at destination below.
                if spec["kind"] == "notion":
                    if payload:
                        body = _notion_text(raw, spec["page_id"]).encode()
                    else:
                        parsed = BeautifulSoup(raw, "html.parser")
                        for tag in parsed(["script", "style"]):
                            tag.decompose()
                        text = parsed.get_text("\n", strip=True)
                        if "VEN-2026-08-26-R2" not in text:
                            raise InputError("Public HTML is a JavaScript shell, not the disclosed brief")
                        body = text.encode()
                else:
                    body = raw
                source["content_type"] = headers.get("Content-Type") or "application/octet-stream"
                source["version_metadata"] = {"http_headers": headers, "transport": "public HTTPS read", "source_identity": spec.get("file_id", spec.get("page_id", spec["url"]))}
                source["retrieved_at"] = now()
                source["attempts"].append({"method": "HTTPS POST read" if payload else "HTTPS GET", "url": url, "observed_at": observed, "completed_at": now(), "http_status": status, "status": "retrieved", "response_sha256": digest(raw)})
                break
            except Exception as error:
                body = None
                source["attempts"].append({"method": "HTTPS read", "url": url, "observed_at": observed, "completed_at": now(), "status": "unavailable", "error": type(error).__name__ + ":" + str(error)[:200]})
    if body is not None:
        destination.write_bytes(body)
        source.update({"content_hash": digest(body), "local_reference": "deliverables/snapshots/evidence/" + destination.name, "retrieval_status": "retrieved"})
        try:
            if spec["kind"] == "xlsx":
                book = load_workbook(destination, read_only=True, data_only=True)
                source["version_metadata"]["tabs"] = book.sheetnames
                book.close()
            elif spec["kind"] == "pdf":
                reader = PdfReader(io.BytesIO(body))
                source["version_metadata"]["page_count"] = len(reader.pages)
                reviewed = digest(body) == "sha256:" + visual_review["sha256"]
                source["visual_review_verified_by_hash"] = reviewed
                if reviewed:
                    source["observations"].append({"id": "OBS-E-FLOOR", "summary": visual_review["observation"], "locator": {"kind": "page-region", "value": visual_review["locator"]}, "review_provenance": visual_review["recorded_in"]})
            elif spec["kind"] == "html":
                if "f1710485827023.pdf" not in body.decode(errors="replace"):
                    raise InputError("Required4F PDF not linked in current accessibility index")
                source["observations"].append({"id": "OBS-E-ACCESS", "summary": "Index links exact required4F PDF", "locator": {"kind": "section", "value": "TICC4F Accessible Facilities Floor Plan link"}})
            elif spec["kind"] == "notion":
                source["version_metadata"]["declared_versions"] = sorted(set(re.findall(r"(?:EVENT-BRIEF|VEN|PROG)-\d{4}-\d{2}-\d{2}(?:-R\d+)?", body.decode())))
        except Exception as error:
            source["retrieval_status"] = "invalid"
            source["interpretation_error"] = type(error).__name__ + ":" + str(error)[:200]
    write_json(evidence / (name + "-capture.json"), source)
    return source
