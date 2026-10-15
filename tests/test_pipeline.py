"""Independent business, artifact, and recovery checks; all runs use temporary roots.

The source bundle is the actual captured exercise evidence. Mutations below are
explicit synthetic test changes; they never assert that a remote source changed.
Run with ``python -m unittest discover -s tests -v`` after installing requirements.
"""

from __future__ import annotations

import csv
import hashlib
import importlib.util
import json
import re
import shutil
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
from icalendar import Calendar
from jsonschema import Draft202012Validator, FormatChecker


REPOSITORY = Path(__file__).resolve().parents[1]
MODULE_PATH = REPOSITORY / "event-planning-coordination-brief/scripts/pipeline.py"
SOURCE_EVIDENCE = REPOSITORY / "deliverables/history/prep-20261014T213226Z/snapshots/evidence"
SCHEMA_HASH = "787e6ac3e2008458b311c2ab2e287a3cafb68aaa82d49a489775864ba32221ee"
ARTIFACTS = (
    "vendor-comparison.csv", "event-plan.md", "event-calendar.ics",
    "draft-communications.md",
)
TAIPEI = ZoneInfo("Asia/Taipei")
STAGES = (
    "scope-and-approval-gates", "source-capture", "constraint-model",
    "planning-baseline", "option-generation", "feasibility-testing",
    "decision-and-approval", "draft-propagation", "publication-validation",
)
FIELDS = {
    "attendee-signals": ("signal_id", "source", "segment", "signal", "value", "unit", "observed_at", "confidence"),
    "budget": ("category", "planned_amount_twd", "approval_limit_twd", "committed_amount_twd", "owner", "record_version"),
    "calendar-constraints": ("constraint_id", "owner", "constraint_type", "start_at", "end_at", "priority", "notes", "record_version"),
    "vendor-register": ("quote_id", "vendor", "category", "option", "quote_amount_twd", "capacity", "available_date", "valid_until", "status", "notes"),
}
KEYS = {"attendee-signals": "signal_id", "budget": "category", "calendar-constraints": "constraint_id", "vendor-register": "quote_id"}


def sha(path):
    return "sha256:" + hashlib.sha256(Path(path).read_bytes()).hexdigest()


def load_pipeline():
    sys.path.insert(0, str(MODULE_PATH.parent))
    spec = importlib.util.spec_from_file_location("event_pipeline_under_test", MODULE_PATH)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def quote_ids(row):
    return set(re.findall(r"Q-\d{3}", str(row.get("quote_ids", ""))))


class PipelineChecks(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pipeline = load_pipeline()
        cls.schema = json.loads((REPOSITORY / "snapshot.schema.json").read_text())
        cls.validator = Draft202012Validator(cls.schema, format_checker=FormatChecker())

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="project-c-independent-")
        self.addCleanup(self.temporary.cleanup)
        self.base = Path(self.temporary.name)
        self.root = self.base / "project"
        self.root.mkdir()
        shutil.copy2(REPOSITORY / "snapshot.schema.json", self.root / "snapshot.schema.json")
        shutil.copytree(REPOSITORY / "interviews", self.root / "interviews")
        self.bundle = self.base / "observed-source-test-bundle"
        shutil.copytree(SOURCE_EVIDENCE, self.bundle)

    def run_pipeline(self, **kwargs):
        return self.pipeline.run(self.root, source_bundle=self.bundle, **kwargs)

    def snapshots(self):
        result = []
        for number, stage in enumerate(STAGES, 1):
            path = self.root / "deliverables/snapshots" / f"{number:02d}-{stage}.json"
            self.assertTrue(path.is_file(), str(path))
            result.append((path, json.loads(path.read_text())))
        return result

    def comparison(self):
        with (self.root / "deliverables/vendor-comparison.csv").open(newline="") as handle:
            return list(csv.DictReader(handle))

    def base_options(self):
        return {
            next(iter(quote_ids(row) & {"Q-003", "Q-004"})): row
            for row in self.comparison()
            if "Q-001" in quote_ids(row)
            and "Q-006" in quote_ids(row)
            and "Q-008" not in quote_ids(row)
            and quote_ids(row) & {"Q-003", "Q-004"}
        }

    def mutate(self, table, key, field, value):
        path = self.bundle / f"{table}.xlsx"
        book = openpyxl.load_workbook(path)
        sheet = book.active
        headers = {cell.value: cell.column for cell in sheet[1]}
        target = next(row for row in sheet.iter_rows(min_row=2) if row[headers[KEYS[table]] - 1].value == key)
        sheet.cell(target[0].row, headers[field], value)
        book.save(path)
        self.update_bundle_hash(path)

    def reset_tables(self):
        for table in FIELDS:
            path = self.bundle / f"{table}.xlsx"
            shutil.copy2(SOURCE_EVIDENCE / path.name, path)
            self.update_bundle_hash(path)

    def remove_record(self, table, key):
        path = self.bundle / f"{table}.xlsx"
        book = openpyxl.load_workbook(path)
        sheet = book.active
        headers = {cell.value: cell.column for cell in sheet[1]}
        target = next(row for row in sheet.iter_rows(min_row=2) if row[headers[KEYS[table]] - 1].value == key)
        sheet.delete_rows(target[0].row)
        book.save(path)
        self.update_bundle_hash(path)

    def update_bundle_hash(self, path):
        registry_path = self.bundle / "source-registry.json"
        registry = json.loads(registry_path.read_text())
        for source in registry["sources"]:
            if Path(source.get("local_reference") or "").name == path.name:
                source["content_hash"] = sha(path)
                source.setdefault("version_metadata", {})["synthetic_test_mutation"] = True
        registry_path.write_text(json.dumps(registry, indent=2) + "\n")

    def assert_chain(self):
        snapshots = self.snapshots()
        run_ids = {data["run_id"] for _, data in snapshots}
        self.assertEqual(len(run_ids), 1)
        self.assertEqual(len({data["snapshot_id"] for _, data in snapshots}), 9)
        produced = set()
        for index, (path, data) in enumerate(snapshots):
            self.validator.validate(data)
            self.assertEqual(data["stage"], STAGES[index])
            self.assertEqual(data["sequence"], index + 1)
            if index == 0:
                self.assertIsNone(data["predecessor"])
            else:
                previous_path, previous = snapshots[index - 1]
                self.assertGreaterEqual(datetime.fromisoformat(data["created_at"].replace("Z", "+00:00")), datetime.fromisoformat(previous["created_at"].replace("Z", "+00:00")))
                self.assertEqual(data["predecessor"]["snapshot_id"], previous["snapshot_id"])
                self.assertEqual(data["predecessor"]["sha256"], sha(previous_path))
                self.assertTrue(set(data["consumed_record_ids"]) <= produced)
            produced.update(data["produced_record_ids"])
        for number in (7, 8):
            data = snapshots[number][1]
            records = data["state"]["artifact_drafts" if number == 7 else "artifacts"]
            self.assertEqual({Path(record["path"]).name for record in records}, set(ARTIFACTS))
            for record in records:
                path = self.root / record["path"]
                self.assertTrue(path.is_file())
                self.assertEqual(record["sha256"], sha(path))
        self.assertEqual(hashlib.sha256((self.root / "snapshot.schema.json").read_bytes()).hexdigest(), SCHEMA_HASH)
        return snapshots

    def test_baseline_artifacts_chain_calendar_and_actual_amounts(self):
        # Read the separately captured CSV representation, independently of the
        # pipeline's XLSX parser and business-model implementation.
        def captured_rows(name):
            with (self.bundle / (name + ".csv")).open(newline="") as handle:
                return list(csv.DictReader(handle))
        attendance = captured_rows("attendee-signals")
        independent_people = sum(int(row["value"]) for row in attendance if row["signal_id"] in {"SIG-001", "SIG-002", "SIG-003", "SIG-004", "SIG-005"})
        independent_registered = sum(int(row["value"]) for row in attendance if row["signal_id"] in {"SIG-001", "SIG-002"})
        self.assertEqual(independent_people, 245)
        self.assertEqual(independent_registered, 180)
        budget_rows = captured_rows("budget")
        self.assertEqual(sum(int(row["planned_amount_twd"]) for row in budget_rows), 510000)
        commitments = sum(int(row["committed_amount_twd"]) for row in budget_rows)
        self.assertEqual(commitments, 0)
        prices = {row["quote_id"]: int(row["quote_amount_twd"]) for row in captured_rows("vendor-register")}
        common_cost = sum(prices[key] for key in ("Q-001", "Q-005", "Q-006", "Q-009", "Q-010")) + 25000 + 40000 + commitments
        self.assertEqual(common_cost + prices["Q-003"], 494000)
        self.assertEqual(common_cost + prices["Q-004"], 512000)
        self.run_pipeline()
        snapshots = self.assert_chain()
        self.assertEqual(snapshots[-1][1]["state"]["publication_status"], "validated")
        options = self.base_options()
        self.assertEqual(set(options), {"Q-003", "Q-004"})
        self.assertEqual(int(options["Q-003"]["cost_twd"]), common_cost + prices["Q-003"])
        self.assertEqual(int(options["Q-004"]["cost_twd"]), common_cost + prices["Q-004"])
        for row in options.values():
            self.assertEqual(int(row["planning_people"]), independent_people)
            self.assertEqual(row["approval_status"], "pending")
            self.assertEqual(row["feasibility"], "feasible")
        rejected_stream = [row for row in self.comparison() if "Q-008" in quote_ids(row)]
        self.assertTrue(rejected_stream)
        self.assertTrue(all(row["feasibility"] == "infeasible" for row in rejected_stream))
        self.assertTrue(any(int(row["cost_twd"]) == 549000 for row in rejected_stream))
        stream_reasons = json.dumps(rejected_stream).lower()
        for word in ("network", "production", "ceiling"):
            self.assertIn(word, stream_reasons)
        for quote, reason in (("Q-002", "keynote"), ("Q-007", "hearing")):
            rejected = [row for row in self.comparison() if quote in quote_ids(row)]
            self.assertTrue(rejected, quote)
            self.assertTrue(all(row["feasibility"] == "infeasible" for row in rejected))
            self.assertIn(reason, json.dumps(rejected).lower())
        self.assertEqual(set().union(*(quote_ids(row) for row in self.comparison())), {f"Q-{number:03d}" for number in range(1, 11)})
        text = (self.root / "deliverables/event-plan.md").read_text()
        for expected in ("245", "144", "40,000", "26,000", "8,000"):
            self.assertIn(expected, text)
        self.assertIn("UNSENT DRAFT", (self.root / "deliverables/draft-communications.md").read_text())
        calendar = Calendar.from_ical((self.root / "deliverables/event-calendar.ics").read_bytes())
        self.assertEqual(str(calendar["VERSION"]), "2.0")
        self.assertEqual(str(calendar["METHOD"]), "PUBLISH")
        events = list(calendar.walk("VEVENT"))
        zone_ids = {str(zone["TZID"]) for zone in calendar.walk("VTIMEZONE")}
        self.assertTrue(events)
        self.assertEqual(len({str(event["UID"]) for event in events}), len(events))
        for event in events:
            self.assertEqual(str(event["STATUS"]), "TENTATIVE")
            self.assertIsNone(event.get("ATTENDEE"))
            self.assertIn("DTSTAMP", event)
            # RFC5545 section3.8.2.2: VEVENT DTEND is strictly after DTSTART.
            self.assertGreater(event["DTEND"].dt, event["DTSTART"].dt)
            for field in ("DTSTART", "DTEND"):
                self.assertIsNotNone(event[field].dt.tzinfo)
                tzid = event[field].params.get("TZID")
                if tzid:
                    self.assertIn(str(tzid), zone_ids)
        full_calendar = calendar.to_ical().decode()
        for number in range(1, 9):
            self.assertIn(f"CAL-{number:03d}", full_calendar)
        deadlines = [item for item in calendar.walk("VTODO") if str(item.get("X-SOURCE-RECORD", "")) == "CAL-005"]
        self.assertEqual(len(deadlines), 1)
        self.assertEqual(str(deadlines[0]["STATUS"]), "NEEDS-ACTION")
        self.assertEqual(deadlines[0]["DUE"].dt.astimezone(TAIPEI).isoformat(), "2026-09-04T17:00:00+08:00")
        self.assertIsNotNone(deadlines[0].get("UID"))
        self.assertIsNotNone(deadlines[0].get("DTSTAMP"))
        demo_events = [event for event in events if re.search(r"DEMO-\d{2}", str(event.get("SUMMARY", "")))]
        self.assertGreaterEqual(len(demo_events), 12)
        demo_minutes = sum((event["DTEND"].dt - event["DTSTART"].dt).total_seconds() / 60 for event in demo_events)
        self.assertEqual(demo_minutes, 156)
        actual_demo_ids = [re.search(r"DEMO-\d{2}", str(event["SUMMARY"])).group() for event in demo_events]
        self.assertEqual(sorted(actual_demo_ids), [f"DEMO-{number:02d}" for number in range(1, 13)])
        for event in demo_events:
            self.assertEqual(event["DTSTART"].dt.astimezone(TAIPEI).date().isoformat(), "2026-10-17")
            self.assertGreaterEqual(event["DTSTART"].dt.astimezone(TAIPEI).strftime("%H:%M"), "09:30")
            self.assertLessEqual(event["DTEND"].dt.astimezone(TAIPEI).strftime("%H:%M"), "17:00")
        # Programme blocks have a dedicated marker; source windows are envelopes,
        # not activities, and should not create false overlap failures.
        programme_events = [event for event in events if str(event.get("X-RECORD-KIND", "")).lower() == "programme"]
        self.assertTrue(programme_events)
        self.assertGreaterEqual(sum("break" in str(event.get("SUMMARY", "")).lower() for event in programme_events), 2)
        ordered = sorted(programme_events, key=lambda event: event["DTSTART"].dt)
        for before, after in zip(ordered, ordered[1:]):
            self.assertLessEqual(before["DTEND"].dt, after["DTSTART"].dt)
        approvals = snapshots[6][1]["state"]["approval_requirements"]
        self.assertTrue(approvals)
        self.assertFalse(any(record["status"] == "approved" for record in approvals))
        self.assertFalse(any("facilitator" in record["summary"].lower() for record in approvals))

    def test_reordered_headers_records_and_extra_columns_combined(self):
        self.run_pipeline()
        expected = {(row["option_id"], row["cost_twd"], row["planning_people"], row["feasibility"]) for row in self.comparison()}
        for table in FIELDS:
            path = self.bundle / f"{table}.xlsx"
            book = openpyxl.load_workbook(path)
            sheet = book.active
            rows = list(sheet.values)
            width = len(rows[0])
            order = list(reversed(range(width)))
            rearranged = [["unrelated_extra"] + [row[index] for index in order] for row in [rows[0]] + list(reversed(rows[1:]))]
            sheet.delete_rows(1, sheet.max_row)
            for row in rearranged:
                sheet.append(row)
            book.save(path)
            self.update_bundle_hash(path)
        self.run_pipeline()
        actual = {(row["option_id"], row["cost_twd"], row["planning_people"], row["feasibility"]) for row in self.comparison()}
        self.assertEqual(expected, actual)
        self.assert_chain()

    def test_missing_and_duplicate_identity_headers_are_rejected(self):
        for mutation in ("missing", "duplicate-header", "duplicate-id"):
            with self.subTest(mutation=mutation):
                path = self.base / f"{mutation}.xlsx"
                book = openpyxl.Workbook()
                sheet = book.active
                if mutation == "missing":
                    sheet.append(["value"])
                    sheet.append([1])
                elif mutation == "duplicate-header":
                    sheet.append(["signal_id", "signal_id", "value"])
                    sheet.append(["A", "B", 1])
                else:
                    sheet.append(["signal_id", "value"])
                    sheet.append(["A", 1])
                    sheet.append(["A", 2])
                book.save(path)
                with self.assertRaises((ValueError, KeyError, RuntimeError)):
                    self.pipeline.parse_table(path, ("signal_id", "value"), "signal_id")

    def test_incompatible_needs_unit_and_ambiguous_calendar_type_are_rejected(self):
        tables = {name: self.pipeline.parse_table(self.bundle / f"{name}.xlsx", fields, KEYS[name]) for name, fields in FIELDS.items()}
        policy = json.loads((REPOSITORY / "event-planning-coordination-brief/references/project.json").read_text())["policy"]
        brief = (self.bundle / "event-brief-and-venue-note.txt").read_text()
        wheelchair = next(row for row in tables["attendee-signals"] if row["signal_id"] == "SIG-006")
        wheelchair["unit"] = "meals"
        with self.assertRaises(ValueError):
            self.pipeline.build_model(tables, policy, brief)
        wheelchair["unit"] = "people"
        duplicated_type = dict(tables["calendar-constraints"][0], constraint_id="CAL-009")
        tables["calendar-constraints"].append(duplicated_type)
        with self.assertRaises(ValueError):
            self.pipeline.build_model(tables, policy, brief)

    def test_attendance_261_invalidates_smaller_package_only(self):
        self.mutate("attendee-signals", "SIG-005", "value", 51)
        self.run_pipeline()
        options = self.base_options()
        self.assertEqual(int(options["Q-003"]["planning_people"]), 261)
        self.assertEqual(options["Q-003"]["feasibility"], "infeasible")
        self.assertIn(options["Q-004"]["feasibility"], ("feasible", "conditional"))
        plan = (self.root / "deliverables/event-plan.md").read_text()
        self.assertIn("261", plan)
        self.assertIn("51", plan)
        self.assert_chain()

    def test_attendance_301_invalidates_both_packages(self):
        self.mutate("attendee-signals", "SIG-005", "value", 91)
        self.run_pipeline()
        options = self.base_options()
        self.assertTrue(all(int(row["planning_people"]) == 301 for row in options.values()))
        self.assertTrue(all(row["feasibility"] == "infeasible" for row in options.values()))
        self.assert_chain()

    def test_commitments_reduce_headroom_and_total_ceiling_is_hard(self):
        self.mutate("budget", "communications", "committed_amount_twd", 27000)
        self.run_pipeline()
        options = self.base_options()
        self.assertEqual(int(options["Q-003"]["cost_twd"]), 521000)
        self.assertEqual(int(options["Q-004"]["cost_twd"]), 539000)
        self.assertTrue(all(row["feasibility"] == "infeasible" for row in options.values()))
        self.assert_chain()

    def test_allocation_variance_differs_from_approval_limit_breach(self):
        # Baseline security28000 is above allocation25000 but below limit35000.
        self.run_pipeline()
        baseline = self.base_options()
        self.assertEqual(baseline["Q-003"]["feasibility"], "feasible")
        self.mutate("vendor-register", "Q-009", "quote_amount_twd", 35000)
        self.run_pipeline()
        at_limit = self.base_options()
        self.assertEqual(int(at_limit["Q-003"]["cost_twd"]), 501000)
        self.assertEqual(at_limit["Q-003"]["feasibility"], "feasible")
        self.mutate("vendor-register", "Q-009", "quote_amount_twd", 36000)
        self.run_pipeline()
        options = self.base_options()
        self.assertEqual(int(options["Q-003"]["cost_twd"]), 502000)
        self.assertEqual(options["Q-003"]["feasibility"], "conditional")
        combined = json.dumps(options["Q-003"]).lower()
        self.assertIn("security", combined)
        self.assertTrue(any(word in combined for word in ("approval", "limit", "threshold")))

    def test_expired_quote_is_evaluated_at_exercise_clock(self):
        self.mutate("vendor-register", "Q-003", "valid_until", datetime(2026, 8, 25))
        self.run_pipeline()
        options = self.base_options()
        self.assertEqual(options["Q-003"]["feasibility"], "infeasible")
        # Q-004 has already expired at the real October retrieval date, but is
        # valid on the required August26 simulated business clock.
        self.assertEqual(options["Q-004"]["feasibility"], "feasible")

    def test_changed_keynote_conflict_with_brief_is_preserved(self):
        self.mutate("calendar-constraints", "CAL-004", "start_at", "2026-10-17T10:30:00+08:00")
        self.mutate("calendar-constraints", "CAL-004", "end_at", "2026-10-17T11:30:00+08:00")
        self.run_pipeline()
        self.assertFalse(any(row["feasibility"] == "feasible" for row in self.base_options().values()))
        snapshots = self.assert_chain()
        conflicts = snapshots[2][1]["state"]["conflicts"]
        self.assertTrue(conflicts)
        self.assertIn("keynote", json.dumps(conflicts).lower())

    def test_each_required_calendar_meaning_is_held_when_missing(self):
        # In particular, deleting a walkthrough/rehearsal/deadline/fallback
        # record must not make those prerequisites silently disappear.
        for number in range(1, 9):
            with self.subTest(missing=f"CAL-{number:03d}"):
                self.reset_tables()
                self.remove_record("calendar-constraints", f"CAL-{number:03d}")
                result = self.run_pipeline()
                self.assertIsNone(result["selected_option"])
                self.assertIn(result["status"], ("partial", "blocked"))
                self.assertTrue(all(row["feasibility"] in ("conditional", "unverified") for row in self.base_options().values()))
                snapshots = self.assert_chain()
                self.assertTrue(any(snapshot["unresolved"] for _, snapshot in snapshots))

    def test_stable_source_ids_cannot_silently_change_business_meaning(self):
        # Paired swaps retain the complete sets of populations/needs/calendar
        # meanings and avoid merely triggering generic duplicate detection.
        cases = (
            ("population-signal", "attendee-signals", "E-ATTENDEES", (("SIG-001", "signal", "expected_attendees"),)),
            ("population-segment", "attendee-signals", "E-ATTENDEES", (("SIG-001", "segment", "partners"), ("SIG-002", "segment", "fellows"))),
            ("overlapping-need", "attendee-signals", "E-ATTENDEES", (("SIG-006", "signal", "mobility_companions"), ("SIG-007", "signal", "wheelchair_spaces"))),
            ("calendar-window-role", "calendar-constraints", "E-CALENDAR", (("CAL-001", "constraint_type", "fallback_event_window"), ("CAL-002", "constraint_type", "preferred_event_window"))),
            ("calendar-prerequisite-role", "calendar-constraints", "E-CALENDAR", (("CAL-006", "constraint_type", "rehearsal_window"), ("CAL-007", "constraint_type", "accessibility_walkthrough"))),
            ("quote-category", "vendor-register", "E-VENDORS", (("Q-005", "category", "production"), ("Q-006", "category", "accessibility"))),
        )
        for name, table, source_id, changes in cases:
            with self.subTest(case=name):
                self.reset_tables()
                for identity, field, value in changes:
                    self.mutate(table, identity, field, value)
                result = self.run_pipeline()
                self.assertIsNone(result["selected_option"])
                snapshots = self.assert_chain()
                source = next(source for source in snapshots[1][1]["state"]["sources"] if source["id"] == source_id)
                self.assertEqual(source["retrieval_status"], "invalid")
                self.assertTrue(source.get("interpretation_error"))

    def test_invalid_required_number_is_not_guessed(self):
        self.mutate("attendee-signals", "SIG-001", "value", "unknown")
        self.run_pipeline()
        text = (self.root / "deliverables/event-plan.md").read_text().lower()
        self.assertTrue(any(word in text for word in ("invalid", "blocked", "unresolved")))
        rows = self.comparison()
        self.assertFalse(any(row["feasibility"] == "feasible" for row in rows))

    def test_missing_floor_plan_withholds_media_claims_and_records_owner(self):
        (self.bundle / "official-4f-floor-plan.pdf").unlink()
        self.run_pipeline()
        snapshots = self.assert_chain()
        source = next(source for source in snapshots[1][1]["state"]["sources"] if source["id"] == "E-FLOOR")
        self.assertIn(source["retrieval_status"], ("unavailable", "invalid", "unverified"))
        self.assertTrue(any(snapshot["unresolved"] for _, snapshot in snapshots))
        self.assertNotEqual(snapshots[-1][1]["status"], "complete")

    def test_missing_vendor_source_never_reuses_prior_success_as_current(self):
        self.run_pipeline()
        old_run = self.snapshots()[0][1]["run_id"]
        (self.bundle / "vendor-register.xlsx").unlink()
        self.run_pipeline()
        snapshots = self.assert_chain()
        source = next(source for source in snapshots[1][1]["state"]["sources"] if source["id"] == "E-VENDORS")
        self.assertIn(source["retrieval_status"], ("unavailable", "invalid", "unverified"))
        self.assertNotEqual(snapshots[0][1]["run_id"], old_run)
        self.assertFalse(any(row["feasibility"] == "feasible" for row in self.comparison()))
        self.assertTrue((self.root / "deliverables/history" / old_run).is_dir())
        self.assertTrue(any(snapshot["unresolved"] for _, snapshot in snapshots))

    def test_unconfirmed_readiness_stays_conditional(self):
        path = self.bundle / "event-brief-and-venue-note.txt"
        text = path.read_text()
        start = text.index("Exercise readiness confirmation")
        end = text.index("Exercise programme confirmation")
        path.write_text(text[:start] + text[end:])
        self.update_bundle_hash(path)
        self.run_pipeline()
        self.assertTrue(all(row["feasibility"] in ("conditional", "unverified", "infeasible") for row in self.base_options().values()))

    def test_corrupt_and_missing_drafts_recover_into_new_preserved_run(self):
        self.run_pipeline()
        first = self.assert_chain()
        first_run = first[0][1]["run_id"]
        first_ids = {snapshot["snapshot_id"] for _, snapshot in first}
        broken = self.root / "deliverables/event-plan.md"
        broken.write_text("CORRUPTED ACTUAL OUTPUT\n")
        (self.root / "deliverables/event-calendar.ics").unlink()
        self.assertFalse(self.pipeline.validate_outputs(self.root)["valid"])
        self.run_pipeline()
        second = self.assert_chain()
        self.assertNotEqual(second[0][1]["run_id"], first_run)
        self.assertEqual(second[0][1]["state"]["changed_decision_ids"], [])
        self.assertFalse(first_ids & {snapshot["snapshot_id"] for _, snapshot in second})
        self.assertNotIn("CORRUPTED ACTUAL OUTPUT", broken.read_text())
        history = self.root / "deliverables/history" / first_run
        self.assertTrue(history.is_dir())
        self.assertTrue(any(path.is_file() and b"CORRUPTED ACTUAL OUTPUT" in path.read_bytes() for path in history.rglob("event-plan.md")))
        Calendar.from_ical((self.root / "deliverables/event-calendar.ics").read_bytes())
        self.assertTrue(self.pipeline.validate_outputs(self.root)["valid"])
        # The corrupted bytes must be diagnosed from actual files, rather than
        # accepting cached successful metadata from the previous run.
        self.assertIn("CORRUPTED ACTUAL OUTPUT", "\n".join(path.read_text(errors="replace") for path in history.rglob("event-plan.md")))

    def test_actual_calendar_duplicate_demo_is_detected(self):
        self.run_pipeline()
        path = self.root / "deliverables/event-calendar.ics"
        calendar = Calendar.from_ical(path.read_bytes())
        second = next(event for event in calendar.walk("VEVENT") if "DEMO-02" in str(event.get("SUMMARY", "")))
        second["SUMMARY"] = "DEMO-01 presentation and changeover"
        path.write_bytes(calendar.to_ical())
        self.assertFalse(self.pipeline.validate_outputs(self.root)["valid"])

    def test_corrupt_predecessor_is_preserved_diagnosed_and_repaired(self):
        self.run_pipeline()
        first = self.snapshots()
        previous_run = first[0][1]["run_id"]
        target, snapshot = first[2]
        snapshot["state"]["assumptions"][0]["summary"] = "CORRUPTED PREDECESSOR TEST BYTES"
        target.write_text(json.dumps(snapshot))
        self.assertNotEqual(sha(target), first[3][1]["predecessor"]["sha256"])
        self.run_pipeline()
        repaired = self.assert_chain()
        self.assertNotEqual(repaired[0][1]["run_id"], previous_run)
        history = self.root / "deliverables/history" / previous_run
        manifest = json.loads((history / "history-manifest.json").read_text())
        self.assertTrue(manifest["observed_anomalies"])
        retained = next(history.rglob("03-constraint-model.json"))
        self.assertIn("CORRUPTED PREDECESSOR TEST BYTES", retained.read_text())

    def test_malformed_prior_provenance_is_preserved_without_invented_decision_change(self):
        for relative in ("snapshots/01-scope-and-approval-gates.json", "snapshots/evidence/source-registry.json"):
            with self.subTest(malformed=relative):
                self.run_pipeline()
                first_run = self.snapshots()[0][1]["run_id"]
                bad_bytes = b"{invalid"
                path = self.root / "deliverables" / relative
                path.write_bytes(bad_bytes)
                result = self.run_pipeline()
                self.assertEqual(result["status"], "complete")
                repaired = self.assert_chain()
                self.assertNotEqual(repaired[0][1]["run_id"], first_run)
                self.assertEqual(repaired[0][1]["state"]["changed_decision_ids"], [])
                self.assertTrue(repaired[0][1]["state"]["prior_output_anomalies"])
                history = self.root / "deliverables/history" / first_run
                self.assertEqual((history / relative).read_bytes(), bad_bytes)
                self.assertTrue(json.loads((history / "history-manifest.json").read_text())["observed_anomalies"])

    def rejected_run(self):
        try:
            result = self.run_pipeline()
        except (ValueError, RuntimeError, OSError):
            return
        self.assertEqual(result["status"], "failed", result)

    def test_history_run_id_traversal_is_rejected_before_outside_write(self):
        self.run_pipeline()
        metadata = self.root / "deliverables/run.json"
        current = json.loads(metadata.read_text())
        current["run_id"] = "../../../escape-from-history"
        metadata.write_text(json.dumps(current))
        outside = self.base / "escape-from-history"
        self.assertFalse(outside.exists())
        self.rejected_run()
        self.assertFalse(outside.exists())

    def test_deliverables_symlink_is_rejected_before_outside_write(self):
        outside = self.base / "outside-deliverables"
        outside.mkdir()
        sentinel = outside / "sentinel.txt"
        sentinel.write_text("UNCHANGED\n")
        (self.root / "deliverables").symlink_to(outside, target_is_directory=True)
        self.rejected_run()
        self.assertEqual(sentinel.read_text(), "UNCHANGED\n")
        self.assertEqual([path.name for path in outside.iterdir()], ["sentinel.txt"])

    def test_history_symlink_is_rejected_before_outside_write(self):
        self.run_pipeline()
        outside = self.base / "outside-history"
        outside.mkdir()
        (self.root / "deliverables/history").symlink_to(outside, target_is_directory=True)
        self.rejected_run()
        self.assertEqual(list(outside.iterdir()), [])


if __name__ == "__main__":
    unittest.main(verbosity=2)
